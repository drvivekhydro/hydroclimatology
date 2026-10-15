"""Builds lec23_worked_examples.xlsx (live formulas) from the tiny CSV inputs next to this script."""
from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

here = Path(__file__).parent
ld = lambda n: np.loadtxt(here / n, delimiter=",", ndmin=2)
BLUE = Font(name="Calibri", color="0000FF")
BLUEB = Font(name="Calibri", color="0000FF", bold=True)
BLACK = Font(name="Calibri", color="000000")
BOLD = Font(name="Calibri", bold=True)
TITLE = Font(name="Calibri", bold=True, size=14)
HEAD = PatternFill("solid", fgColor="DDEBF7")
SEC = PatternFill("solid", fgColor="FFF2CC")
wb = Workbook()


def put(ws, ref, v, kind="f", fmt=None, bold=False):
    c = ws[ref]
    c.value = v
    if kind == "i":
        c.font = BLUEB if bold else BLUE
    else:
        c.font = BOLD if bold else BLACK
    if fmt:
        c.number_format = fmt
    return c


def label(ws, ref, text, section=False):
    c = ws[ref]
    c.value = text
    c.font = BOLD
    if section:
        c.fill = SEC


def head(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = BOLD
    c.fill = HEAD
    c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")


def note(ws, ref, text):
    cm = Comment(text, "Lecture 23")
    cm.width = 380
    cm.height = 140
    ws[ref].comment = cm


def matrix(ws, top, left, arr, kind="i", fmt="0.0"):
    for r, row in enumerate(arr):
        for c, v in enumerate(row):
            put(ws, f"{L(left + c)}{top + r}", float(v), kind, fmt)


def heat(ws, rng, rev=False):
    lo, hi = ("F8696B", "63BE7B") if rev else ("FFFFFF", "2F75B5")
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=lo, end_type="max", end_color=hi))


# ============================== README ==============================
ws = wb.active
ws.title = "README"
rows = [
    ("Lecture 23: Spatial downscaling. Companion workbook (CE524 Applied Hydroclimatology, IIT Mandi)", TITLE),
    ("", None),
    ("Purpose: hand-checkable worked examples of the lecture, with live formulas: (1) bilinear interpolation, (2) BCSD from a 2x2 coarse to a 4x4 fine grid, (3) constructed analogues with a ridge term.", None),
    ("ALL NUMBERS ARE SYNTHETIC. They were invented for teaching; none is an observation or a model output. To use real data, replace the blue cells on the sheet Data (e.g. IMD gridded rainfall for the fine fields, a CMIP6 or NEX-GDDP-CMIP6 field averaged to the coarse cells).", None),
    ("", None),
    ("Colour code: BLUE font = input (type over it). BLACK font = formula (do not overwrite). Yellow bars = step headings. Blue shading = column headings. Hover over a heading with a red corner to read a note about the formula.", None),
    ("", None),
    ("Sheets", BOLD),
    ("Data: every input of the three examples in one place.", None),
    ("Bilinear: 4x4 coarse grid, six target points; weights from the formulas, two-step versus four-weight value, nearest neighbour.", None),
    ("BCSD_2x2: factor = coarse GCM / coarse climatology, interpolated factor, rescaling by the fine climatology (rainfall, multiplicative); anomaly version for temperature (additive).", None),
    ("BCCA: 8 library days, RMSE ranking, the N=3 best analogues, weights by LINEST, by the normal equations and by ridge (MMULT/MINVERSE), downscaled 4x4 field.", None),
    ("", None),
    ("How to use: change a blue input (for example the ridge lambda on Data) and watch the formulas follow. Set lambda = 0 on Data: the ridge weights become the plain least-squares weights.", None),
    ("Only widely supported functions are used (AVERAGE, SUMPRODUCT, INDEX, MATCH, RANK, IF, MIN, MAX, INT, MMULT, MINVERSE, TRANSPOSE, LINEST, SQRT, SUMXMY2). The matrix formulas are legacy array formulas (they work in Excel and Google Sheets).", None),
    ("The file was written by a script (openpyxl) and carries no stored results; Excel computes them when the file is opened (press F9 if cells look empty). The same numbers are reproduced in lecture-23-spatial-downscaling.ipynb.", None),
    ("Units: precipitation mm/day, temperature degC. Grid indices: row 0 is the first row listed, column 0 the first column listed.", None),
]
for i, (t, f) in enumerate(rows, 1):
    ws[f"A{i}"] = t
    ws[f"A{i}"].font = f or Font(name="Calibri")
    ws[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140
put(ws, "A18", "Example of an input (blue):")
put(ws, "B18", 1.0, "i")
put(ws, "A19", "Example of a formula (black):")
put(ws, "B19", "=B18*2", "f")
ws.column_dimensions["B"].width = 10

# ============================== Data ==============================
d = wb.create_sheet("Data")
d["A1"] = "Data: synthetic inputs for the three worked examples (blue = input)"
d["A1"].font = TITLE
d["A2"] = "Same numbers as the CSV files in data/downscaling/lec23/ (bilinear_*, bcsd_*, ca_*)."
label(d, "A4", "A. Coarse grid z (mm/day) at nodes x = 0..3 (columns), y = 0..3 (rows)", True)
head(d, "A5", "y \\ x")
for k in range(4):
    put(d, f"{L(2 + k)}5", k, "i", "0", True)
    put(d, f"A{6 + k}", k, "i", "0", True)
matrix(d, 6, 2, ld("bilinear_coarse_4x4.csv"), "i", "0.0")
heat(d, "B6:E9")
label(d, "G4", "A. Target points (x, y in coarse-cell units)", True)
for c, t in zip("GHI", ["id", "x", "y"]):
    head(d, f"{c}5", t)
pts = ld("bilinear_points_xy.csv")
for i, (x, y) in enumerate(pts):
    put(d, f"G{6 + i}", i + 1, "i", "0")
    put(d, f"H{6 + i}", float(x), "i", "0.00")
    put(d, f"I{6 + i}", float(y), "i", "0.00")
label(d, "A12", "B. BCSD inputs: precipitation (mm/day), 2x2 coarse and 4x4 fine; temperature (degC)", True)
label(d, "A13", "B1 Coarse climatology of observations")
matrix(d, 14, 2, ld("bcsd_coarse_clim_2x2.csv"))
label(d, "A17", "B2 Bias-corrected coarse GCM, today")
matrix(d, 18, 2, ld("bcsd_gcm_bc_today_2x2.csv"))
label(d, "A21", "B3 Fine climatology of observations")
matrix(d, 22, 2, ld("bcsd_fine_clim_4x4.csv"))
heat(d, "B22:E25")
label(d, "G13", "B4 Coarse clim. T (degC)")
matrix(d, 14, 8, ld("bcsd_T_coarse_clim_2x2.csv"))
label(d, "G17", "B5 BC-GCM T today (degC)")
matrix(d, 18, 8, ld("bcsd_T_gcm_bc_today_2x2.csv"))
label(d, "G21", "B6 Fine clim. T (degC)")
matrix(d, 22, 8, ld("bcsd_T_fine_clim_4x4.csv"))
heat(d, "H22:K25", rev=True)
label(d, "A28", "C. Constructed analogues: library of 8 days, fine 4x4 field per day (mm/day); fine cell id = 4*row + col + 1", True)
head(d, "A29", "day")
for k in range(16):
    head(d, f"{L(2 + k)}29", k + 1)
lib = ld("ca_library_fine_8x16.csv")
for i in range(8):
    put(d, f"A{30 + i}", i + 1, "i", "0")
matrix(d, 30, 2, lib, "i", "0.0")
heat(d, "B30:Q37")
label(d, "A39", "Target coarse field, cells (1,1) (1,2) (2,1) (2,2)")
matrix(d, 39, 2, ld("ca_target_coarse_1x4.csv"))
label(d, "A40", "Truth fine field (only to check)")
matrix(d, 40, 2, ld("ca_target_truth_fine_1x16.csv"))
label(d, "A42", "Ridge lambda (mm2/day2)")
put(d, "B42", 5.0, "i", "0.00")
note(d, "B42", "Ridge parameter. 0 gives plain least squares. Larger values shrink the weights toward zero.")
for c in range(2, 18):
    d.column_dimensions[L(c)].width = 7
d.column_dimensions["A"].width = 46
d.column_dimensions["G"].width = 24
d.freeze_panes = "B4"

# ============================== Bilinear ==============================
b = wb.create_sheet("Bilinear")
b["A1"] = "Bilinear interpolation: linear in x, then linear in y; equal to the area-weighted mean of the 4 corners"
b["A1"].font = TITLE
b["A2"] = "x runs along columns, y along rows. (i0, j0) is the lower-left node of the cell that holds the point; tx, ty in [0,1] are the fractions of the way across it."
hd = ["id", "x", "y", "i0", "j0", "tx", "ty", "w00", "w10", "w01", "w11", "sum w", "z00", "z10", "z01", "z11",
      "z bottom (x first)", "z top (x first)", "z two-step", "z four-weight", "two-step minus four-weight",
      "nearest neighbour", "nearest minus bilinear"]
for k, h in enumerate(hd):
    head(b, f"{L(1 + k)}4", h)
b.row_dimensions[4].height = 45
G = "Data!$B$6:$E$9"
for r in range(5, 11):
    q = r + 1
    f = {
        "A": f"=Data!G{q}", "B": f"=Data!H{q}", "C": f"=Data!I{q}",
        "D": f"=MIN(INT(B{r}),COLUMNS({G})-2)", "E": f"=MIN(INT(C{r}),ROWS({G})-2)",
        "F": f"=B{r}-D{r}", "G": f"=C{r}-E{r}",
        "H": f"=(1-F{r})*(1-G{r})", "I": f"=F{r}*(1-G{r})", "J": f"=(1-F{r})*G{r}", "K": f"=F{r}*G{r}",
        "L": f"=SUM(H{r}:K{r})",
        "M": f"=INDEX({G},E{r}+1,D{r}+1)", "N": f"=INDEX({G},E{r}+1,D{r}+2)",
        "O": f"=INDEX({G},E{r}+2,D{r}+1)", "P": f"=INDEX({G},E{r}+2,D{r}+2)",
        "Q": f"=(1-F{r})*M{r}+F{r}*N{r}", "R": f"=(1-F{r})*O{r}+F{r}*P{r}",
        "S": f"=(1-G{r})*Q{r}+G{r}*R{r}", "T": f"=SUMPRODUCT(H{r}:K{r},M{r}:P{r})", "U": f"=S{r}-T{r}",
        "V": f"=INDEX({G},INT(C{r}+0.5)+1,INT(B{r}+0.5)+1)", "W": f"=V{r}-T{r}"}
    for c, v in f.items():
        fmt = "0" if c in "ADE" else ("0.00" if c in "BCFG" else ("0.0000" if c in "HIJKLU" else "0.000"))
        put(b, f"{c}{r}", v, "f", fmt)
label(b, "A12", "Interpolation adds no information: the results stay inside the range of the coarse values (smoothness, no new extremes)", True)
for i, (t, fm) in enumerate([("Smallest coarse value", f"=MIN({G})"), ("Largest coarse value", f"=MAX({G})"),
                             ("Smallest interpolated value", "=MIN(T5:T10)"), ("Largest interpolated value", "=MAX(T5:T10)")]):
    put(b, f"A{13 + i}", t, "f", None, True)
    put(b, f"E{13 + i}", fm, "f", "0.000")
note(b, "D4", "i0 = column index of the left node. INT takes the whole part; MIN keeps the point inside the last cell when x equals the last node.")
note(b, "F4", "tx = x - i0: the fraction of the way from the left node to the right node. Same for ty.")
note(b, "H4", "w00 = (1-tx)(1-ty) is the area of the rectangle OPPOSITE the node, so the closer the point, the larger the weight. The four weights add to 1 (column L).")
note(b, "Q4", "Step 1 of the derivation: linear interpolation along x on the lower edge (y = j0).")
note(b, "R4", "Linear interpolation along x on the upper edge (y = j0+1).")
note(b, "S4", "Step 2: linear interpolation along y between the two edge values.")
note(b, "T4", "Same number as S, computed as the weighted sum of the four corners (SUMPRODUCT of the weights and the corner values).")
note(b, "V4", "Value of the nearest node: INT(x+0.5) rounds to the nearest whole number.")
for c in range(2, 24):
    b.column_dimensions[L(c)].width = 11
b.column_dimensions["A"].width = 28
b.freeze_panes = "B5"

# ============================== BCSD_2x2 ==============================
s = wb.create_sheet("BCSD_2x2")
s["A1"] = "BCSD, tiny example: 2x2 coarse to 4x4 fine (quantile mapping assumed done: the coarse GCM field is already bias corrected)"
s["A1"].font = TITLE
s["A2"] = "Rainfall: MULTIPLICATIVE scaling (factor = today / climatology). Temperature (lower part): ADDITIVE scaling (anomaly = today - climatology)."
label(s, "A4", "Step 0. Inputs (links to Data)", True)
label(s, "A5", "Coarse clim. Ybar (mm/day)")
label(s, "E5", "BC-GCM coarse today y (mm/day)")
label(s, "I5", "Fine clim. xbar (mm/day)")
for r in range(2):
    for c in range(2):
        put(s, f"{L(2 + c)}{6 + r}", f"=Data!{L(2 + c)}{14 + r}", "f", "0.00")
        put(s, f"{L(6 + c)}{6 + r}", f"=Data!{L(2 + c)}{18 + r}", "f", "0.00")
for r in range(4):
    for c in range(4):
        put(s, f"{L(10 + c)}{6 + r}", f"=Data!{L(2 + c)}{22 + r}", "f", "0.00")
label(s, "A9", "Check: block means of xbar (= Ybar?)")
for r in range(2):
    for c in range(2):
        put(s, f"{L(2 + c)}{10 + r}", f"=AVERAGE({L(10 + 2 * c)}{6 + 2 * r}:{L(11 + 2 * c)}{7 + 2 * r})", "f", "0.00")
label(s, "A13", "Step 1. Factor R = y / Ybar (>1: wetter than usual)", True)
for r in range(2):
    for c in range(2):
        put(s, f"{L(2 + c)}{14 + r}", f"={L(6 + c)}{6 + r}/{L(2 + c)}{6 + r}", "f", "0.000")
heat(s, "B14:C15")
label(s, "A17", "Step 2. Interpolation weights along one axis (fine cell centres vs coarse cell centres)", True)
label(s, "A18", "Refinement factor f (fine cells per coarse cell)")
put(s, "C18", 2, "i", "0")
label(s, "A19", "First coarse centre c0 = (f-1)/2 (fine-cell units)")
put(s, "C19", "=(C18-1)/2", "f", "0.00")
label(s, "A20", "Second coarse centre c1 = c0 + f")
put(s, "C20", "=C19+C18", "f", "0.00")
label(s, "A21", "Fine cell centre index xi")
label(s, "A22", "Weight t = clip((xi - c0)/(c1 - c0), 0, 1)")
for c in range(4):
    put(s, f"{L(2 + c)}21", c, "i", "0")
    put(s, f"{L(2 + c)}22", f"=MIN(1,MAX(0,({L(2 + c)}21-$C$19)/($C$20-$C$19)))", "f", "0.00")
note(s, "A22", "t is the fraction of the way from the first to the second coarse centre. Fine cells outside the two centres (xi = 0 and 3) are clamped to t = 0 or 1, i.e. they take the nearest coarse value (no extrapolation).")
label(s, "A24", "Step 3. Interpolated factor R_f (along x with t, then along y with t_y)", True)
head(s, "F24", "t_y (row)")
for r in range(4):
    put(s, f"A{25 + r}", r, "i", "0")
    put(s, f"F{25 + r}", f"=INDEX($B$22:$E$22,A{25 + r}+1)", "f", "0.00")
    for c in range(4):
        cl = L(2 + c)
        put(s, f"{cl}{25 + r}", f"=(1-$F{25 + r})*((1-{cl}$22)*$B$14+{cl}$22*$C$14)+$F{25 + r}*((1-{cl}$22)*$B$15+{cl}$22*$C$15)", "f", "0.000")
heat(s, "B25:E28")
label(s, "A30", "Step 4. Downscaled rainfall P = R_f x xbar (mm/day)", True)
for r in range(4):
    for c in range(4):
        put(s, f"{L(2 + c)}{31 + r}", f"={L(2 + c)}{25 + r}*{L(10 + c)}{6 + r}", "f", "0.00")
heat(s, "B31:E34")
label(s, "G30", "For comparison: plain bilinear interpolation of y (no climatology)")
for r in range(4):
    for c in range(4):
        cl = L(2 + c)
        put(s, f"{L(8 + c)}{31 + r}", f"=(1-$F{25 + r})*((1-{cl}$22)*$F$6+{cl}$22*$G$6)+$F{25 + r}*((1-{cl}$22)*$F$7+{cl}$22*$G$7)", "f", "0.00")
heat(s, "H31:K34")
label(s, "A36", "Step 5. Check: block means of P, and minus y", True)
for r in range(2):
    for c in range(2):
        put(s, f"{L(2 + c)}{37 + r}", f"=AVERAGE({L(2 + 2 * c)}{31 + 2 * r}:{L(3 + 2 * c)}{32 + 2 * r})", "f", "0.00")
        put(s, f"{L(6 + c)}{37 + r}", f"={L(2 + c)}{37 + r}-{L(6 + c)}{6 + r}", "f", "0.00")
label(s, "E36", "difference:")
note(s, "A36", "BCSD as usually described does not force the block means to equal the coarse value; the difference is small when the factor field is smooth. Some implementations add a final rescaling to enforce it.")
label(s, "A40", "Temperature, additive version", True)
label(s, "A41", "Coarse clim. Tbar (degC)")
label(s, "E41", "BC-GCM T today (degC)")
label(s, "I41", "Fine clim. T (degC)")
for r in range(2):
    for c in range(2):
        put(s, f"{L(2 + c)}{42 + r}", f"=Data!{L(8 + c)}{14 + r}", "f", "0.0")
        put(s, f"{L(6 + c)}{42 + r}", f"=Data!{L(8 + c)}{18 + r}", "f", "0.0")
for r in range(4):
    for c in range(4):
        put(s, f"{L(10 + c)}{42 + r}", f"=Data!{L(8 + c)}{22 + r}", "f", "0.0")
label(s, "A47", "Anomaly A = y - Tbar (K)")
for r in range(2):
    for c in range(2):
        put(s, f"{L(2 + c)}{48 + r}", f"={L(6 + c)}{42 + r}-{L(2 + c)}{42 + r}", "f", "0.0")
label(s, "A51", "Interpolated anomaly A_f (same weights as Step 3)")
for r in range(4):
    for c in range(4):
        cl = L(2 + c)
        put(s, f"{cl}{52 + r}", f"=(1-$F{25 + r})*((1-{cl}$22)*$B$48+{cl}$22*$C$48)+$F{25 + r}*((1-{cl}$22)*$B$49+{cl}$22*$C$49)", "f", "0.00")
label(s, "A57", "Downscaled T = fine clim + A_f (degC)")
for r in range(4):
    for c in range(4):
        put(s, f"{L(2 + c)}{58 + r}", f"={L(10 + c)}{42 + r}+{L(2 + c)}{52 + r}", "f", "0.00")
heat(s, "B58:E61", rev=True)
label(s, "G57", "Block means of T, minus y")
for r in range(2):
    for c in range(2):
        put(s, f"{L(8 + c)}{58 + r}", f"=AVERAGE({L(2 + 2 * c)}{58 + 2 * r}:{L(3 + 2 * c)}{59 + 2 * r})-{L(6 + c)}{42 + r}", "f", "0.00")
for c in range(2, 14):
    s.column_dimensions[L(c)].width = 11
s.column_dimensions["A"].width = 48
s.freeze_panes = "B4"

# ============================== BCCA ==============================
a = wb.create_sheet("BCCA")
a["A1"] = "Constructed analogues: today's coarse field as a combination of the coarse fields of N=3 similar days; reuse the weights on their fine fields"
a["A1"].font = TITLE
a["A2"] = "q = 4 coarse cells, p = 16 fine cells, library of H = 8 days. No intercept: weights by plain least squares (LINEST) or ridge (MMULT/MINVERSE)."
label(a, "A4", "Step 1. Coarse fields of the library days (mean of each 2x2 block of the fine field) and RMSE to today's target", True)
for k, h in enumerate(["day", "cell 1", "cell 2", "cell 3", "cell 4", "RMSE to target", "rank (1 = most similar)"]):
    head(a, f"{L(1 + k)}5", h)
a.row_dimensions[5].height = 32
for i in range(8):
    r = 6 + i
    dr = 30 + i
    put(a, f"A{r}", f"=Data!A{dr}", "f", "0")
    for I in range(2):
        for J in range(2):
            ids = [4 * (2 * I) + 2 * J + 1 + o for o in (0, 1, 4, 5)]
            cells = ",".join(f"Data!{L(1 + i2)}{dr}" for i2 in ids)
            put(a, f"{L(2 + 2 * I + J)}{r}", f"=AVERAGE({cells})", "f", "0.00")
    put(a, f"F{r}", f"=SQRT(SUMPRODUCT((B{r}:E{r}-Data!$B$39:$E$39)^2)/COLUMNS(B{r}:E{r}))", "f", "0.000")
    put(a, f"G{r}", f"=RANK(F{r},$F$6:$F$13,1)", "f", "0")
label(a, "A15", "Target y (BC-GCM)")
for c in range(4):
    put(a, f"{L(2 + c)}15", f"=Data!{L(2 + c)}39", "f", "0.00")
note(a, "F5", "RMSE = sqrt(mean over the 4 coarse cells of (library - target)^2): how different the day's coarse pattern is from today's.")
note(a, "G5", "RANK(...,1) ranks ascending: 1 is the smallest RMSE, i.e. the best analogue.")
label(a, "A17", "Step 2. The N = 3 best analogue days", True)
label(a, "A18", "k")
label(a, "A19", "day index")
label(a, "A20", "RMSE")
for k in range(3):
    cl = L(2 + k)
    put(a, f"{cl}18", k + 1, "i", "0")
    put(a, f"{cl}19", f"=MATCH({cl}18,$G$6:$G$13,0)", "f", "0")
    put(a, f"{cl}20", f"=INDEX($F$6:$F$13,{cl}19)", "f", "0.000")
label(a, "A22", "Step 3. Y (q x N): coarse fields of the chosen days as columns; y: target", True)
for c, h in zip("ABCDE", ["cell", "day k=1", "day k=2", "day k=3", "y (target)"]):
    head(a, f"{c}23", h)
for i in range(4):
    r = 24 + i
    put(a, f"A{r}", i + 1, "i", "0")
    for k in range(3):
        cl = L(2 + k)
        put(a, f"{cl}{r}", f"=INDEX($B$6:$E$13,{cl}$19,$A{r})", "f", "0.00")
    put(a, f"E{r}", f"=INDEX($B$15:$E$15,A{r})", "f", "0.00")
label(a, "A29", "Y'Y (N x N)")
label(a, "F29", "Y'y (N x 1)")
a["B30"] = ArrayFormula("B30:D32", "=MMULT(TRANSPOSE(B24:D27),B24:D27)")
a["F30"] = ArrayFormula("F30:F32", "=MMULT(TRANSPOSE(B24:D27),E24:E27)")
for r in range(30, 33):
    for c in "BCDF":
        a[f"{c}{r}"].number_format = "0.00"
        a[f"{c}{r}"].font = BLACK
note(a, "B29", "Matrix (legacy array) formula over B30:D32: Y-transpose times Y. Entry (m,n) is the dot product of the coarse fields of analogue m and analogue n.")
label(a, "A34", "Ridge: G = Y'Y + lambda I   (lambda from Data)", True)
label(a, "A35", "lambda")
put(a, "B35", "=Data!B42", "f", "0.00")
for r in range(3):
    for c in range(3):
        src = f"{L(2 + c)}{30 + r}"
        put(a, f"{L(2 + c)}{36 + r}", f"={src}+$B$35" if r == c else f"={src}", "f", "0.00")
label(a, "F35", "G inverse")
a["F36"] = ArrayFormula("F36:H38", "=MINVERSE(B36:D38)")
for r in range(36, 39):
    for c in "FGH":
        a[f"{c}{r}"].number_format = "0.0000"
        a[f"{c}{r}"].font = BLACK
label(a, "A40", "Step 4. Weights", True)
for k, h in enumerate(["k", "plain LS (LINEST)", "plain LS (normal equations)", "ridge (MMULT, MINVERSE)"]):
    head(a, f"{L(1 + k)}41", h)
a.row_dimensions[41].height = 32
a["D42"] = ArrayFormula("D42:D44", "=MMULT(F36:H38,F30:F32)")
a["C42"] = ArrayFormula("C42:C44", "=MMULT(MINVERSE(B30:D32),F30:F32)")
for k in range(3):
    r = 42 + k
    put(a, f"A{r}", k + 1, "i", "0")
    put(a, f"B{r}", f"=INDEX(LINEST($E$24:$E$27,$B$24:$D$27,FALSE,FALSE),1,COLUMNS($B$24:$D$24)+1-A{r})", "f", "0.0000")
    for c in "CD":
        a[f"{c}{r}"].number_format = "0.0000"
        a[f"{c}{r}"].font = BLACK
label(a, "A45", "sum of weights")
for c in "BCD":
    put(a, f"{c}45", f"=SUM({c}42:{c}44)", "f", "0.0000")
note(a, "B41", "LINEST returns the coefficients in REVERSE order of the X columns, and const=FALSE forces no intercept, so INDEX(...,1,N+1-k) picks the weight of analogue k.")
note(a, "C41", "Same answer from the normal equations (Y'Y)^-1 Y'y. With lambda = 0 the ridge column equals these two columns; set Data!B42 = 0 to check.")
note(a, "D41", "Ridge: (Y'Y + lambda I)^-1 Y'y. Keeps the inverse well defined and the weights bounded when analogues are nearly collinear or N exceeds the number of coarse cells.")
label(a, "A47", "Step 5. Same weights on the FINE fields of the chosen days", True)
for k, h in enumerate(["fine cell", "row", "col", "fine, day k=1", "fine, day k=2", "fine, day k=3", "plain LS", "ridge", "truth", "error ridge", "error plain"]):
    head(a, f"{L(1 + k)}48", h)
a.row_dimensions[48].height = 32
for i in range(16):
    r = 49 + i
    put(a, f"A{r}", i + 1, "i", "0")
    put(a, f"B{r}", f"=INT((A{r}-1)/4)", "f", "0")
    put(a, f"C{r}", f"=MOD(A{r}-1,4)", "f", "0")
    for k in range(3):
        put(a, f"{L(4 + k)}{r}", f"=INDEX(Data!$B$30:$Q$37,{L(2 + k)}$19,$A{r})", "f", "0.0")
    put(a, f"G{r}", f"=D{r}*$B$42+E{r}*$B$43+F{r}*$B$44", "f", "0.00")
    put(a, f"H{r}", f"=D{r}*$D$42+E{r}*$D$43+F{r}*$D$44", "f", "0.00")
    put(a, f"I{r}", f"=INDEX(Data!$B$40:$Q$40,A{r})", "f", "0.0")
    put(a, f"J{r}", f"=H{r}-I{r}", "f", "0.00")
    put(a, f"K{r}", f"=G{r}-I{r}", "f", "0.00")
label(a, "A66", "RMSE vs truth (mm/day): best single analogue (D), plain LS (G), ridge (H)", True)
put(a, "D66", "=SQRT(SUMXMY2(D49:D64,I49:I64)/COUNT(I49:I64))", "f", "0.000")
put(a, "G66", "=SQRT(SUMXMY2(G49:G64,I49:I64)/COUNT(I49:I64))", "f", "0.000")
put(a, "H66", "=SQRT(SUMXMY2(H49:H64,I49:I64)/COUNT(I49:I64))", "f", "0.000")
label(a, "A67", "Mean of ridge field (H) vs mean of target y (I)")
put(a, "H67", "=AVERAGE(H49:H64)", "f", "0.000")
put(a, "I67", "=AVERAGE(B15:E15)", "f", "0.000")
label(a, "A69", "Maps (row x col): ridge estimate (left) and truth (right)", True)
for c in range(4):
    put(a, f"{L(2 + c)}70", c, "i", "0")
    put(a, f"{L(8 + c)}70", c, "i", "0")
for r in range(4):
    put(a, f"A{71 + r}", r, "i", "0")
    put(a, f"G{71 + r}", r, "i", "0")
    for c in range(4):
        put(a, f"{L(2 + c)}{71 + r}", f"=INDEX($H$49:$H$64,4*$A{71 + r}+{L(2 + c)}$70+1)", "f", "0.0")
        put(a, f"{L(8 + c)}{71 + r}", f"=INDEX($I$49:$I$64,4*$G{71 + r}+{L(8 + c)}$70+1)", "f", "0.0")
heat(a, "B71:E74")
heat(a, "H71:K74")
for c in range(2, 12):
    a.column_dimensions[L(c)].width = 13
a.column_dimensions["A"].width = 38
ch = BarChart()
ch.type = "col"
ch.title = "Analogue weights"
ch.y_axis.title = "weight"
ch.x_axis.title = "analogue k"
ch.add_data(Reference(a, min_col=2, max_col=4, min_row=41, max_row=44), titles_from_data=True)
ch.set_categories(Reference(a, min_col=1, min_row=42, max_row=44))
ch.height = 7
ch.width = 14
a.add_chart(ch, "M4")
a.freeze_panes = "B4"

wb.calculation.fullCalcOnLoad = True
wb.save(here / "lec23_worked_examples.xlsx")
print("saved")
